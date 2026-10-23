"""Regression checks for the existing native report + LibreOffice PDF path.

Run from OntiCards_Api: python -m unittest discover -s tests -p test_governance_pdf_export.py -v
The real conversion check runs when LibreOffice Writer is installed.
"""
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from controllers.governance import libreoffice_exporter as exporter


def sample_report():
    report = SimpleNamespace(
        report_name='合成治理报告', execution_time=datetime(2026, 10, 22, 12),
        datasource_id='synthetic-source', scope_tables=['maintenance_orders'],
        quality_score=66.7, grade='一般',
        details={'relation_discovery': {'relationships_count': 1, 'tables_count': 2,
                 'cards_count': 2, 'cross_source_count': 0,
                 'statistics': {'foreign_key': {'count': 1}}}},
    )
    row = SimpleNamespace(
        rule_name='订单金额非空', table_name='maintenance_orders', column_name='amount',
        rule_type='null_check', severity='warning', status='failed', failed_rate=33.3,
        failed_count=1, total_count=3, passed_count=2,
        failed_samples=[{'order_id': 2, 'violated_condition': 'amount IS NOT NULL',
                         'condition_expr': 'amount IS NOT NULL'}],
    )
    return report, [row]


class GovernancePdfExportTest(unittest.TestCase):
    def test_native_document_keeps_report_sections_and_fits_print_page(self):
        report, rows = sample_report()
        with tempfile.TemporaryDirectory() as directory:
            filename = str(Path(directory) / 'report.docx')
            exporter._build_docx_with_python(report, rows, filename)
            doc = Document(filename)
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            for required in ['maintenance_orders', 'amount IS NOT NULL', '改进建议',
                             '表关系发现结果', 'foreign_key', '影响行数: 1', '1/3']:
                self.assertIn(required, text)
            section = doc.sections[0]
            self.assertGreater(section.page_width, section.page_height)
            available = section.page_width - section.left_margin - section.right_margin
            self.assertLessEqual(sum(c.width for c in doc.tables[0].columns), available + 1000)

    def test_missing_converter_reports_failure_without_creating_pdf(self):
        report, rows = sample_report()
        with tempfile.TemporaryDirectory() as directory:
            instance = exporter.LibreOfficeExporter(export_path=directory)
            with patch.object(instance, '_check_soffice_available', return_value=None):
                with self.assertRaisesRegex(RuntimeError, 'LibreOffice Writer'):
                    instance.export(report, rows, fmt='pdf')
            self.assertEqual(list(Path(directory).glob('*.pdf')), [])

    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice Writer is required')
    def test_real_pdf_conversion_avoids_html_import_and_produces_file(self):
        report, rows = sample_report()
        with tempfile.TemporaryDirectory() as directory:
            instance = exporter.LibreOfficeExporter(export_path=directory)
            with patch.object(exporter, '_build_report_html', side_effect=AssertionError('HTML import regressed')):
                result = instance.export(report, rows, fmt='pdf')
            data = Path(result['file_path']).read_bytes()
            self.assertTrue(data.startswith(b'%PDF-'))
            self.assertGreater(len(data), 1000)
            self.assertEqual(result['format'], 'pdf')
            self.assertEqual(result['file_size'], len(data))
            if os.environ.get('ONTICARDS_PDF_TEST_OUTPUT'):
                target = Path(os.environ['ONTICARDS_PDF_TEST_OUTPUT'])
                target.mkdir(parents=True, exist_ok=True)
                shutil.copy2(result['file_path'], target / 'failed-rule-and-relation.pdf')


if __name__ == '__main__':
    unittest.main()
